#!/usr/bin/env python3
"""Convertit le classeur « Colombe3_Suivi_Design » en données initiales de l'application.

Usage : python3 scripts/excel_vers_donnees.py <classeur.xlsx> web/donnees-initiales.js

- Feuille « Suivi Détaillé » -> lots et postes (budget, prestataire, observations).
- Feuille « Historique »     -> paiements, rattachés chacun à un poste.

Dans l'application, le montant payé d'un poste est la somme de ses paiements.
Le script vérifie donc, poste par poste, que cette somme est égale à la colonne
« MONTANT PAYÉ » du classeur ; si un écart subsiste, un paiement
« Reprise du classeur Excel » est ajouté pour le combler et signalé à l'écran.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime

import openpyxl

# Libellés de l'historique qui ne correspondent pas mot pour mot à un poste.
ALIAS = {
    "Ferronnerie — portails & portions (solde)": "Ferronnerie",
    "Fermeture regards + Fosse septique": "Fermeture regards (×9) + Fosse septique",
    "Ramassage gravats": "Ramassage gravats de construction",
    "Sacs de plâtre (80 sacs)": "Achat sacs de plâtre (80 sacs)",
    "Sacs de plâtre (25 sacs)": "Achat sacs de plâtre (25 sacs)",
    "Ciment (2 tonnes)": "Achat ciment (2 tonnes)",
    "Houdis": "Achat houdis",
    "Achat tuyaux orange — filerie ext.": "Achat tuyaux orange — filerie ext. (clôture & jardin)",
    "Achat ciment — divers travaux": "Achat ciment (nouveau) — divers travaux",
    "Mise à niveau cours, plateforme garage & cuisine africaine":
        "Mise à niveau — cours, plateforme garage & cuisine africaine",
    "Main d'œuvre maçon — mise à niveau cours/plateforme":
        "Mise à niveau — cours, plateforme garage & cuisine africaine",
    "Aluminium, Inox & Vitrerie — avance": "Aluminium, Inox & Vitrerie",
    "Location groupe électrogène": "Location groupe électrogène (01/07/2026)",
    "Achat de disque pour meule": "Achat de disque pour meule (installation portillons)",
    "Menuiserie — portes": "Menuiserie — 8 portes iroko (devis DEVHUB)",
    "Menuiserie — 8 portes iroko": "Menuiserie — 8 portes iroko (devis DEVHUB)",
    "Appareillage électrique (dont DPN tableau électrique)":
        "Appareillage électrique (luminaire, interrupteur, prises, etc.)",
    "Appareillage électrique — câble électrique":
        "Appareillage électrique (luminaire, interrupteur, prises, etc.)",
    "Aménagement jardin — complément main d'œuvre":
        "Aménagement jardin — gazon, pelouse & main d'œuvre",
    "Finition sanitaire — fournitures": "Finition sanitaire — fournitures (devis 30/04/2026)",
}


def cle(texte):
    texte = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()


def date_iso(valeur):
    if valeur is None or valeur == "":
        return ""
    if isinstance(valeur, datetime):
        return valeur.strftime("%Y-%m-%d")
    j, m, a = str(valeur).strip().split("/")
    return f"{a}-{int(m):02d}-{int(j):02d}"


def texte(valeur):
    return "" if valeur is None else str(valeur).strip()


def nombre(valeur):
    return 0 if valeur in (None, "") else round(float(valeur))


def feuille(classeur, mot):
    for ws in classeur.worksheets:
        if mot.lower() in ws.title.lower():
            return ws
    raise SystemExit(f"Feuille « {mot} » introuvable")


def main(source, cible):
    classeur = openpyxl.load_workbook(source)
    suivi = feuille(classeur, "Suivi")
    historique = feuille(classeur, "Historique")

    lots, postes, payes = [], [], {}
    for ligne in suivi.iter_rows(min_row=3, values_only=True):
        nom, date, prestataire, budget, paye, _, _, _, obs, lot = ligne[:10]
        if not nom:
            continue
        entete = re.match(r"▌\s*LOT\s+([A-Z])\s*—", str(nom))
        if entete:
            continue
        m = re.match(r"Lot\s+([A-Z])\s*—\s*(.+)", texte(lot))
        if not m:
            raise SystemExit(f"Lot inconnu pour « {nom} »")
        code, libelle = m.group(1), m.group(2).strip()
        if not any(l["id"] == code for l in lots):
            lots.append({"id": code, "nom": libelle})
        poste = {
            "id": f"p{len(postes) + 1}",
            "lot": code,
            "nom": texte(nom),
            "date": date_iso(date),
            "prestataire": texte(prestataire),
            "budget": nombre(budget),
            "obs": texte(obs),
        }
        postes.append(poste)
        payes[poste["id"]] = nombre(paye)

    par_nom = {cle(p["nom"]): p for p in postes}
    paiements, inconnus = [], []
    for ligne in historique.iter_rows(min_row=2, values_only=True):
        date, nom, prestataire, montant, mode, notes = ligne[:6]
        if not nom or str(date or "").startswith("TOTAL") or str(nom).startswith("TOTAL"):
            continue
        poste = par_nom.get(cle(ALIAS.get(texte(nom), texte(nom))))
        if poste is None:
            inconnus.append(texte(nom))
            continue
        paiements.append({
            "id": f"v{len(paiements) + 1}",
            "posteId": poste["id"],
            "date": date_iso(date),
            "montant": nombre(montant),
            "mode": texte(mode),
            "prestataire": texte(prestataire) or poste["prestataire"],
            "notes": texte(notes),
        })
    if inconnus:
        raise SystemExit("Paiements sans poste correspondant : " + ", ".join(inconnus))

    ecarts = []
    for poste in postes:
        verse = sum(v["montant"] for v in paiements if v["posteId"] == poste["id"])
        ecart = payes[poste["id"]] - verse
        if ecart:
            ecarts.append((poste["nom"], ecart))
            paiements.append({
                "id": f"v{len(paiements) + 1}",
                "posteId": poste["id"],
                "date": poste["date"],
                "montant": ecart,
                "mode": "",
                "prestataire": poste["prestataire"],
                "notes": "Reprise du classeur Excel (montant payé non détaillé dans l'historique)",
            })

    donnees = {
        "version": 1,
        "projet": {"nom": "Colombe 3", "sousTitre": "Suivi budgétaire · Finition duplex"},
        "lots": lots,
        "postes": postes,
        "paiements": paiements,
    }
    with open(cible, "w", encoding="utf-8") as f:
        f.write("// Généré par scripts/excel_vers_donnees.py — ne pas modifier à la main.\n")
        f.write("window.DONNEES_INITIALES = ")
        json.dump(donnees, f, ensure_ascii=False, indent=1)
        f.write(";\n")

    total_budget = sum(p["budget"] for p in postes)
    total_paye = sum(v["montant"] for v in paiements)
    print(f"{len(lots)} lots, {len(postes)} postes, {len(paiements)} paiements")
    print(f"Budget total : {total_budget:,} | Payé : {total_paye:,} | Reste : {total_budget - total_paye:,}")
    for nom, ecart in ecarts:
        print(f"  Écart comblé par une reprise : {nom} -> {ecart:,}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
